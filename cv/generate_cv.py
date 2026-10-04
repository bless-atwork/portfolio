"""Generate Blessing Kwenda's professional CV (DOCX + PDF)."""

from __future__ import annotations

import shutil
import subprocess
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Cm, Pt, RGBColor

OUT_DIR = Path(__file__).resolve().parent
DOWNLOADS = Path.home() / "Downloads"
NAVY = RGBColor(0x1A, 0x2B, 0x4A)
ACCENT = RGBColor(0x2C, 0x3E, 0x5A)
MUTED = RGBColor(0x44, 0x44, 0x44)
BLACK = RGBColor(0x1A, 0x1A, 0x1A)


def set_run_font(run, *, size=10, bold=False, italic=False, color=BLACK, name="Calibri"):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run.font.name = name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), name)
    rFonts.set(qn("w:hAnsi"), name)
    rFonts.set(qn("w:eastAsia"), name)


def set_paragraph_spacing(paragraph, *, before=0, after=4, line=1.08):
    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = line


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1A2B4A")
    pBdr.append(bottom)
    pPr.append(pBdr)


def section_heading(doc, text):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=8, after=3, line=1.0)
    run = p.add_run(text.upper())
    set_run_font(run, size=11, bold=True, color=NAVY)
    add_horizontal_line(p)
    return p


def body_para(doc, text, *, size=10, bold=False, italic=False, before=0, after=4):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=before, after=after, line=1.08)
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold, italic=italic, color=BLACK)
    return p


def bullet(doc, text, *, bold_prefix=None, level=0):
    p = doc.add_paragraph(style="List Bullet")
    set_paragraph_spacing(p, before=0, after=1, line=1.02)
    p.paragraph_format.left_indent = Cm(0.45 + level * 0.35)
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, size=10, bold=True, color=BLACK)
        r2 = p.add_run(text)
        set_run_font(r2, size=10, color=BLACK)
    else:
        run = p.add_run(text)
        set_run_font(run, size=10, color=BLACK)
    return p


def role_line(doc, title, dates):
    p = doc.add_paragraph()
    set_paragraph_spacing(p, before=4, after=1, line=1.0)
    r1 = p.add_run(title)
    set_run_font(r1, size=10.5, bold=True, color=ACCENT)
    tab = p.add_run("\t")
    set_run_font(tab, size=10, color=MUTED)
    r2 = p.add_run(dates)
    set_run_font(r2, size=10, italic=True, color=MUTED)
    tab_stops = p.paragraph_format.tab_stops
    tab_stops.add_tab_stop(Cm(17.5), alignment=2)
    return p


def build_cv() -> Document:
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(1.15)
    section.bottom_margin = Cm(1.15)
    section.left_margin = Cm(1.55)
    section.right_margin = Cm(1.55)
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)

    # --- Header ---
    name = doc.add_paragraph()
    name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(name, before=0, after=2, line=1.0)
    nr = name.add_run("BLESSING KWENDA")
    set_run_font(nr, size=20, bold=True, color=NAVY, name="Calibri")

    headline = doc.add_paragraph()
    headline.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(headline, before=0, after=4, line=1.0)
    hr = headline.add_run("FULL-STACK DEVELOPER  |  DATA SCIENCE  |  SOFTWARE ENGINEERING")
    set_run_font(hr, size=10.5, bold=True, color=ACCENT)

    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(contact, before=0, after=2, line=1.05)
    cr = contact.add_run(
        "Pretoria, South Africa  ·  +27 68 240 6775  ·  blessatwork@gmail.com"
    )
    set_run_font(cr, size=9.5, color=MUTED)

    links = doc.add_paragraph()
    links.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_spacing(links, before=0, after=6, line=1.05)
    lr = links.add_run(
        "github.com/bless-atwork  ·  blessatworkportolio.web.app"
    )
    set_run_font(lr, size=9.5, color=MUTED)

    # --- Professional Summary ---
    section_heading(doc, "Professional Summary")
    body_para(
        doc,
        "Resourceful Full-Stack Developer and Data Science practitioner with a completed "
        "undergraduate Bachelor of Science in Information Technology: Data Science and a "
        "Higher Certificate in Information Systems (Software Development) from Eduvos. "
        "Ships production web "
        "products end-to-end — React/TypeScript front ends, Node.js APIs, Firebase cloud "
        "services, and secure client-facing workflows — while applying Python/R analytics, "
        "NLP, and visualisation where data drives decisions. Currently deepening cybersecurity "
        "competence through the ISC2 pathway. Comfortable owning features across product, "
        "engineering, and operations in real business environments.",
        after=2,
    )

    # --- Areas of Expertise ---
    section_heading(doc, "Areas of Expertise")

    body_para(doc, "Full-Stack & Product Development", size=10.5, bold=True, before=2, after=1)
    bullet(
        doc,
        "React, TypeScript, Vite, Tailwind CSS; modular SPAs, permission-aware navigation, real-time UIs",
    )
    bullet(
        doc,
        "Node.js, Express, REST APIs; Firebase (Auth, Firestore, Storage, Hosting); JWT and MFA (TOTP)",
    )
    bullet(
        doc,
        "SaaS CRM workflows: contacts, pipelines, projects, invoicing, notifications, client portals",
    )
    bullet(
        doc,
        "PDF e-sign / Smart Docs (PDF.js, pdf-lib); email systems (Resend, Nodemailer); RBAC & security rules",
    )

    body_para(doc, "Data Science & Analytics", size=10.5, bold=True, before=4, after=1)
    bullet(
        doc,
        "Python for data science; R analytics; exploratory analysis, data mining, and ETL",
    )
    bullet(
        doc,
        "Natural Language Processing (NLTK); text mining, frequency analysis, and visualisation",
    )
    bullet(
        doc,
        "Time-series analysis and forecasting (SAS Viya); storytelling with data for technical and non-technical audiences",
    )

    body_para(doc, "Software Development Foundations", size=10.5, bold=True, before=4, after=1)
    bullet(doc, "Languages: Java, C#, PHP, Python, C++, JavaScript, HTML, CSS")
    bullet(
        doc,
        "Program design, OOP, UML, software engineering; Linux; Git; database design (SQL, MySQL, SQL Server, MongoDB)",
    )

    body_para(doc, "Cybersecurity Foundations", size=10.5, bold=True, before=4, after=1)
    bullet(
        doc,
        "ISC2 cybersecurity certifications pathway (in progress) — security foundations, "
        "risk awareness, secure practices, and professional ethics",
    )

    body_para(doc, "Professional Skills", size=10.5, bold=True, before=4, after=1)
    bullet(doc, "Strong problem-solving, ownership, and end-to-end delivery")
    bullet(doc, "Clear communication across technical and business stakeholders")
    bullet(doc, "Attention to detail, adaptability, and cross-functional collaboration")

    # --- Professional Expertise / Experience ---
    section_heading(doc, "Professional Expertise")

    role_line(
        doc,
        "Origin CRM — Full-Stack Developer / Product Builder",
        "Current",
    )
    body_para(
        doc,
        "React, TypeScript, Vite, Tailwind · Node.js, Express · Firebase · Resend / Nodemailer · "
        "PDF.js / pdf-lib · Quill · JWT / MFA (TOTP)",
        size=9,
        italic=True,
        after=1,
    )
    body_para(
        doc,
        "Built and shipped a production-ready CRM for sales, operations, and client work — "
        "contacts, pipelines, projects, invoicing, document e-sign, notifications, and "
        "client-facing portals — with a React SPA and a secure Node API on Firebase.",
        size=9.5,
        after=2,
    )
    bullet(
        doc,
        "Designed and developed an end-to-end CRM covering contacts/organizations, leads, "
        "deal pipelines, project boards, activities, and role-based access.",
    )
    bullet(
        doc,
        "Built a React + TypeScript SPA with modular routing, permission-aware navigation, "
        "and real-time dashboard snapshots from Firestore.",
    )
    bullet(
        doc,
        "Implemented a Node.js/Express API with Firebase Admin for authenticated staff flows, "
        "public document/scheduling links, email delivery, and background jobs.",
    )
    bullet(
        doc,
        "Delivered Smart Docs: Word/PDF import, PDF canvas overlays, public signing links, "
        "and high-quality pdf-lib stamped PDF downloads/email attachments.",
    )
    bullet(
        doc,
        "Built invoicing & payments tracking, a client portal, and notification systems "
        "(in-app + email) with user preferences and document-signed alerts.",
    )
    bullet(
        doc,
        "Hardened security with Firebase Auth, MFA (OTP), Storage/Firestore rules, "
        "staff-only APIs, and tokenized public document access.",
    )

    role_line(doc, "Personal Assistant — Jewel Boutique", "Jun 2022 – Present")
    body_para(
        doc,
        "Hartebeespoort, North West, South Africa",
        size=9.5,
        italic=True,
        after=1,
    )
    bullet(
        doc,
        "Provide operational and administrative support in a retail environment, balancing "
        "reliability with concurrent product and technical delivery.",
    )
    bullet(
        doc,
        "Deliver technical assistance including support for the business Google website presence.",
    )
    body_para(
        doc,
        "Reference contact: Sipelile — 062 060 9070",
        size=9,
        italic=True,
        after=2,
    )

    # --- Education ---
    section_heading(doc, "Education")

    role_line(
        doc,
        "Eduvos — Bachelor of Science in Information Technology: Data Science",
        "2022 – 2025",
    )
    body_para(
        doc,
        "Undergraduate degree · Completed  |  Pretoria, South Africa",
        size=9.5,
        italic=True,
        after=1,
    )
    bullet(
        doc,
        "Completed undergraduate Data Science degree covering Python and R analytics, data "
        "mining and administration, NLP with Python, time-series forecasting, microservices "
        "concepts, data visualisation/communication, and research-oriented project work.",
    )

    role_line(
        doc,
        "Eduvos — Higher Certificate in Information Systems: Software Development",
        "Mar 2021 – Jan 2022",
    )
    body_para(
        doc,
        "SAQA ID 120688 · NQF Level 5  |  Pretoria, South Africa",
        size=9.5,
        italic=True,
        after=1,
    )
    bullet(
        doc,
        "Career-focused qualification with practical foundations in programming (Java, C#, PHP, "
        "Python program design), database design and management (MySQL / SQL Server), Linux, "
        "software engineering with UML, and technical project work.",
    )

    # --- Projects ---
    section_heading(doc, "Selected Projects")

    role_line(doc, "Olympic Medal Prediction Model", "Sep 2024 – Present")
    bullet(
        doc,
        "Collaborative predictive modelling using event performance, GDP, GDP per capita, "
        "and population to forecast Olympic outcomes.",
    )

    role_line(doc, "UN Declaration Text Analytics", "May 2023")
    body_para(
        doc,
        "github.com/bless-atwork/data_analytics",
        size=9,
        italic=True,
        after=1,
    )
    bullet(
        doc,
        "Python / Jupyter pipeline for word clouds and frequency analysis of the UN Declaration "
        "of Human Rights — NLP preprocessing and visual storytelling.",
    )

    role_line(doc, "CreditAccess Website", "Sep 2022 – Present")
    bullet(
        doc,
        "Web application built in C#, applying object-oriented programming and web development concepts.",
    )

    # --- Highlights ---
    section_heading(doc, "Highlights & Certifications")
    bullet(
        doc,
        "Block Two Top Achiever — awarded by Eduvos for academic performance.",
        bold_prefix="Academic: ",
    )
    bullet(
        doc,
        "ISC2 Cybersecurity Certifications pathway — advancing cybersecurity careers and "
        "building employer confidence (foundations, risk awareness, secure practices, ethics).",
        bold_prefix="In progress: ",
    )
    bullet(
        doc,
        "Higher Certificate in Information Systems: Software Development; "
        "undergraduate Bachelor of Science in Information Technology: Data Science.",
        bold_prefix="Completed: ",
    )

    # --- Profile links ---
    section_heading(doc, "Profile Links")
    bullet(doc, "GitHub: https://github.com/bless-atwork")
    bullet(doc, "Portfolio: https://blessatworkportolio.web.app")

    # --- Interests ---
    section_heading(doc, "Interests")
    bullet(doc, "Full-stack product building and collaborative technical projects")
    bullet(doc, "Sports, calm community gatherings, landscape photography, and music")

    # --- References ---
    section_heading(doc, "References")
    body_para(
        doc,
        "References available upon request. Ready to contribute to high-impact, product-driven teams.",
        after=0,
    )

    return doc


def export_pdf_via_word(docx_path: Path, pdf_path: Path) -> bool:
    """Convert DOCX to PDF using Microsoft Word COM automation on Windows."""
    try:
        import win32com.client  # type: ignore
    except ImportError:
        return False

    word = None
    try:
        word = win32com.client.Dispatch("Word.Application")
        word.Visible = False
        doc = word.Documents.Open(str(docx_path))
        doc.SaveAs(str(pdf_path), FileFormat=17)
        doc.Close(False)
        return True
    except Exception as exc:  # noqa: BLE001
        print(f"Word COM export failed: {exc}")
        return False
    finally:
        if word is not None:
            try:
                word.Quit()
            except Exception:  # noqa: BLE001
                pass


def write_html_fallback(html_path: Path) -> None:
    """Print-friendly HTML used for PDF export."""
    html = """<!DOCTYPE html>
<html lang="en">
<head>
<meta charset="utf-8" />
<title>Blessing Kwenda — CV</title>
<style>
  @page { size: A4; margin: 1.05cm 1.4cm; }
  body { font-family: Calibri, Arial, sans-serif; color: #1a1a1a; font-size: 9.5pt; line-height: 1.26; max-width: 21cm; margin: 0 auto; padding: 0; }
  h1 { text-align: center; color: #1a2b4a; font-size: 17pt; margin: 0 0 1px; letter-spacing: 0.04em; }
  .headline { text-align: center; color: #2c3e5a; font-weight: 700; font-size: 9.8pt; margin: 0 0 4px; }
  .contact { text-align: center; color: #555; font-size: 8.8pt; margin: 0 0 1px; }
  h2 { color: #1a2b4a; font-size: 10pt; text-transform: uppercase; border-bottom: 1.5px solid #1a2b4a; margin: 8px 0 3px; padding-bottom: 1px; }
  h3 { font-size: 9.6pt; margin: 4px 0 1px; color: #2c3e5a; }
  .meta { color: #666; font-style: italic; font-size: 8.6pt; margin: 0 0 2px; }
  ul { margin: 1px 0 3px 1.0em; padding: 0; }
  li { margin: 1px 0; }
  .row { display: flex; justify-content: space-between; gap: 10px; align-items: baseline; }
  .dates { color: #666; font-style: italic; white-space: nowrap; font-size: 9pt; }
  p { margin: 2px 0 3px; }
</style>
</head>
<body>
  <h1>BLESSING KWENDA</h1>
  <p class="headline">FULL-STACK DEVELOPER  |  DATA SCIENCE  |  SOFTWARE ENGINEERING</p>
  <p class="contact">Pretoria, South Africa · +27 68 240 6775 · blessatwork@gmail.com</p>
  <p class="contact">github.com/bless-atwork · blessatworkportolio.web.app</p>

  <h2>Professional Summary</h2>
  <p>Resourceful Full-Stack Developer and Data Science practitioner with a completed undergraduate Bachelor of Science in Information Technology: Data Science and a Higher Certificate in Information Systems (Software Development) from Eduvos. Ships production web products end-to-end — React/TypeScript front ends, Node.js APIs, Firebase cloud services, and secure client-facing workflows — while applying Python/R analytics, NLP, and visualisation where data drives decisions. Currently deepening cybersecurity competence through the ISC2 pathway. Comfortable owning features across product, engineering, and operations in real business environments.</p>

  <h2>Areas of Expertise</h2>
  <h3>Full-Stack &amp; Product Development</h3>
  <ul>
    <li>React, TypeScript, Vite, Tailwind CSS; modular SPAs, permission-aware navigation, real-time UIs</li>
    <li>Node.js, Express, REST APIs; Firebase (Auth, Firestore, Storage, Hosting); JWT and MFA (TOTP)</li>
    <li>SaaS CRM workflows: contacts, pipelines, projects, invoicing, notifications, client portals</li>
    <li>PDF e-sign / Smart Docs (PDF.js, pdf-lib); email systems (Resend, Nodemailer); RBAC &amp; security rules</li>
  </ul>
  <h3>Data Science &amp; Analytics</h3>
  <ul>
    <li>Python for data science; R analytics; exploratory analysis, data mining, and ETL</li>
    <li>Natural Language Processing (NLTK); text mining, frequency analysis, and visualisation</li>
    <li>Time-series analysis and forecasting (SAS Viya); storytelling with data for technical and non-technical audiences</li>
  </ul>
  <h3>Software Development Foundations</h3>
  <ul>
    <li>Languages: Java, C#, PHP, Python, C++, JavaScript, HTML, CSS</li>
    <li>Program design, OOP, UML, software engineering; Linux; Git; database design (SQL, MySQL, SQL Server, MongoDB)</li>
  </ul>
  <h3>Cybersecurity Foundations</h3>
  <ul>
    <li>ISC2 cybersecurity certifications pathway (in progress) — security foundations, risk awareness, secure practices, and professional ethics</li>
  </ul>
  <h3>Professional Skills</h3>
  <ul>
    <li>Strong problem-solving, ownership, and end-to-end delivery</li>
    <li>Clear communication across technical and business stakeholders</li>
    <li>Attention to detail, adaptability, and cross-functional collaboration</li>
  </ul>

  <h2>Professional Expertise</h2>
  <div class="row"><h3>Origin CRM — Full-Stack Developer / Product Builder</h3><span class="dates">Current</span></div>
  <p class="meta">React, TypeScript, Vite, Tailwind · Node.js, Express · Firebase · Resend / Nodemailer · PDF.js / pdf-lib · Quill · JWT / MFA (TOTP)</p>
  <p>Built and shipped a production-ready CRM for sales, operations, and client work — contacts, pipelines, projects, invoicing, document e-sign, notifications, and client-facing portals — with a React SPA and a secure Node API on Firebase.</p>
  <ul>
    <li>Designed and developed an end-to-end CRM covering contacts/organizations, leads, deal pipelines, project boards, activities, and role-based access.</li>
    <li>Built a React + TypeScript SPA with modular routing, permission-aware navigation, and real-time dashboard snapshots from Firestore.</li>
    <li>Implemented a Node.js/Express API with Firebase Admin for authenticated staff flows, public document/scheduling links, email delivery, and background jobs.</li>
    <li>Delivered Smart Docs: Word/PDF import, PDF canvas overlays, public signing links, and high-quality pdf-lib stamped PDF downloads/email attachments.</li>
    <li>Built invoicing &amp; payments tracking, a client portal, and notification systems (in-app + email) with user preferences and document-signed alerts.</li>
    <li>Hardened security with Firebase Auth, MFA (OTP), Storage/Firestore rules, staff-only APIs, and tokenized public document access.</li>
  </ul>

  <div class="row"><h3>Personal Assistant — Jewel Boutique</h3><span class="dates">Jun 2022 – Present</span></div>
  <p class="meta">Hartebeespoort, North West, South Africa</p>
  <ul>
    <li>Provide operational and administrative support in a retail environment, balancing reliability with concurrent product and technical delivery.</li>
    <li>Deliver technical assistance including support for the business Google website presence.</li>
  </ul>
  <p class="meta">Reference contact: Sipelile — 062 060 9070</p>

  <h2>Education</h2>
  <div class="row"><h3>Eduvos — Bachelor of Science in Information Technology: Data Science</h3><span class="dates">2022 – 2025</span></div>
  <p class="meta">Undergraduate degree · Completed | Pretoria, South Africa</p>
  <ul><li>Completed undergraduate Data Science degree covering Python and R analytics, data mining and administration, NLP with Python, time-series forecasting, microservices concepts, data visualisation/communication, and research-oriented project work.</li></ul>

  <div class="row"><h3>Eduvos — Higher Certificate in Information Systems: Software Development</h3><span class="dates">Mar 2021 – Jan 2022</span></div>
  <p class="meta">SAQA ID 120688 · NQF Level 5 | Pretoria, South Africa</p>
  <ul><li>Career-focused qualification with practical foundations in programming (Java, C#, PHP, Python program design), database design and management (MySQL / SQL Server), Linux, software engineering with UML, and technical project work.</li></ul>

  <h2>Selected Projects</h2>
  <div class="row"><h3>Olympic Medal Prediction Model</h3><span class="dates">Sep 2024 – Present</span></div>
  <ul><li>Collaborative predictive modelling using event performance, GDP, GDP per capita, and population to forecast Olympic outcomes.</li></ul>
  <div class="row"><h3>UN Declaration Text Analytics</h3><span class="dates">May 2023</span></div>
  <p class="meta">github.com/bless-atwork/data_analytics</p>
  <ul><li>Python / Jupyter pipeline for word clouds and frequency analysis of the UN Declaration of Human Rights — NLP preprocessing and visual storytelling.</li></ul>
  <div class="row"><h3>CreditAccess Website</h3><span class="dates">Sep 2022 – Present</span></div>
  <ul><li>Web application built in C#, applying object-oriented programming and web development concepts.</li></ul>

  <h2>Highlights &amp; Certifications</h2>
  <ul>
    <li><strong>Academic:</strong> Block Two Top Achiever — awarded by Eduvos for academic performance.</li>
    <li><strong>In progress:</strong> ISC2 Cybersecurity Certifications pathway — advancing cybersecurity careers and building employer confidence (foundations, risk awareness, secure practices, ethics).</li>
    <li><strong>Completed:</strong> Higher Certificate in Information Systems: Software Development; undergraduate Bachelor of Science in Information Technology: Data Science.</li>
  </ul>

  <h2>Profile Links</h2>
  <ul>
    <li>GitHub: https://github.com/bless-atwork</li>
    <li>Portfolio: https://blessatworkportolio.web.app</li>
  </ul>

  <h2>Interests</h2>
  <ul>
    <li>Full-stack product building and collaborative technical projects</li>
    <li>Sports, calm community gatherings, landscape photography, and music</li>
  </ul>

  <h2>References</h2>
  <p>References available upon request. Ready to contribute to high-impact, product-driven teams.</p>
</body>
</html>
"""
    html_path.write_text(html, encoding="utf-8")


def main():
    docx_path = OUT_DIR / "Blessing_Kwenda_CV.docx"
    pdf_path = OUT_DIR / "Blessing_Kwenda_CV.pdf"
    html_path = OUT_DIR / "Blessing_Kwenda_CV.html"

    doc = build_cv()
    doc.save(str(docx_path))
    print(f"Wrote {docx_path}")

    write_html_fallback(html_path)
    if pdf_path.exists():
        pdf_path.unlink()

    pdf_ok = False
    browsers = [
        r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Microsoft\Edge\Application\msedge.exe",
        r"C:\Program Files\Google\Chrome\Application\chrome.exe",
        r"C:\Program Files (x86)\Google\Chrome\Application\chrome.exe",
    ]

    for browser in browsers:
        if Path(browser).exists():
            try:
                subprocess.check_call(
                    [
                        browser,
                        "--headless=new",
                        "--disable-gpu",
                        "--no-pdf-header-footer",
                        f"--print-to-pdf={pdf_path}",
                        html_path.as_uri(),
                    ],
                    stdout=subprocess.DEVNULL,
                    stderr=subprocess.DEVNULL,
                )
                if pdf_path.exists():
                    pdf_ok = True
                    print(f"Wrote PDF via browser: {pdf_path}")
                    break
            except Exception as exc:  # noqa: BLE001
                print(f"Browser PDF attempt failed ({browser}): {exc}")

    if not pdf_ok:
        pdf_ok = export_pdf_via_word(docx_path, pdf_path)
        if pdf_ok:
            print(f"Wrote PDF via Word: {pdf_path}")
        else:
            print(f"PDF export unavailable — HTML fallback kept at: {html_path}")

    DOWNLOADS.mkdir(parents=True, exist_ok=True)
    shutil.copy2(docx_path, DOWNLOADS / docx_path.name)
    print(f"Copied DOCX to {DOWNLOADS / docx_path.name}")
    if pdf_ok and pdf_path.exists():
        shutil.copy2(pdf_path, DOWNLOADS / pdf_path.name)
        print(f"Copied PDF to {DOWNLOADS / pdf_path.name}")
    elif html_path.exists():
        shutil.copy2(html_path, DOWNLOADS / html_path.name)
        print(f"Copied HTML to {DOWNLOADS / html_path.name}")

    print("Done.")


if __name__ == "__main__":
    main()
